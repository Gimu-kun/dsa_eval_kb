#!/usr/bin/env python3
"""Sinh instance Question / ExpectedAnswer / ExpectedRule và assertion tương ứng từ Excel + mapping ý chính."""
from __future__ import annotations

import json
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT = os.path.dirname(os.path.abspath(__file__))
ONTO = os.path.join(ROOT, "ontology")
XLSX = os.path.join(ROOT, "Ontology.xlsx")

REL_TYPE = "REL_QUESTION_HAS_QUESTION_TYPE"
REL_DIFF = "REL_QUESTION_HAS_DIFFICULTY"
REL_BLOOM = "REL_QUESTION_HAS_BLOOM_LEVEL"
REL_EA = "REL_QUESTION_HAS_EXPECTED_ANSWER"
REL_RUBRIC = "REL_QUESTION_HAS_RUBRICS"
REL_SCORE = "REL_EXPECTED_ANSWER_HAS_SCORING_RULE"
REL_CONCEPT = "REL_EXPECTED_RULE_REQUIRES_CONCEPT"

RUBRICS = [
    ("O_C_RUBRIC_ACCURACY", 0.4),
    ("O_C_RUBRIC_COMPLETENESS", 0.4),
    ("O_C_RUBRIC_LOGICAL_ORDER", 0.2),
]


def R(name, weight, concepts, miss, wrong, rule=None, fn=None):
    item = {
        "name": name,
        "weight": weight,
        "concepts": concepts,
        "miss": miss,
        "wrong": wrong,
    }
    if rule:
        item["domainRuleId"] = rule
    if fn:
        item["functionId"] = fn
    return item


# Ý chính theo từng câu — không phủ chi tiết phụ ngoài ontology.
MODELS = {
    "Q_C1_R_DES_E": [
        R("Thuật toán hữu hạn", 0.35, ["O_C_ALGORITHM", "O_C_FINITENESS"],
          "Thiếu ý thuật toán là dãy hữu hạn bước và phải dừng.",
          "Mô tả hữu hạn sai hoặc nhầm với vòng lặp vô hạn.",
          rule="rul-000004", fn="FUN_IS_TERMINATE"),
        R("Input và output", 0.35, ["O_C_INPUT", "O_C_OUTPUT"],
          "Thiếu dữ liệu đầu vào/đầu ra của thuật toán.",
          "Nhầm input với output hoặc không xác định I/O.",
          rule="rul-000002", fn="FUN_IS_IDENTIFY"),
        R("Giải bài toán xác định", 0.30, ["O_C_PROBLEM_SOLVING", "O_C_DEFINITENESS"],
          "Thiếu ý thuật toán nhằm giải một bài toán với bước xác định.",
          "Không gắn thuật toán với bài toán hoặc tính xác định."),
    ],
    "Q_C1_R_DES_M": [
        R("Big-O là cận trên", 0.40, ["O_C_BIG_O", "O_C_COMPLEXITY"],
          "Thiếu định nghĩa Big-O là cận trên tiệm cận của chi phí.",
          "Nhầm Big-O với cận dưới hoặc chi phí đúng bằng."),
        R("Bỏ hằng số hạng thấp", 0.30, ["O_C_TIME_COMPLEXITY", "O_C_SPACE_COMPLEXITY"],
          "Thiếu ý bỏ hằng số và hạng bậc thấp khi n đủ lớn.",
          "Giữ hằng số như một phần của ký hiệu Big-O."),
        R("Ví dụ bậc phức tạp", 0.30, ["O_C_CONSTANT_TIME", "O_C_LINEAR_TIME", "O_C_QUADRATIC_TIME"],
          "Thiếu ví dụ O(1), O(n) hoặc O(n^2).",
          "Gán sai ý nghĩa các ký hiệu Big-O.",
          fn="FUN_IS_BETTER_NOTATION"),
    ],
    "Q_C1_R_DES_H": [
        R("Tính xác định", 0.20, ["O_C_DEFINITENESS", "O_C_INPUT", "O_C_OUTPUT"],
          "Thiếu đặc trưng xác định (definiteness) và I/O rõ ràng.",
          "Mô tả xác định sai hoặc bỏ I/O.",
          rule="rul-000002", fn="FUN_IS_IDENTIFY"),
        R("Tính hữu hạn", 0.20, ["O_C_FINITENESS"],
          "Thiếu đặc trưng hữu hạn: thuật toán phải dừng.",
          "Nhầm hữu hạn với số bước cố định không phụ thuộc dữ liệu.",
          rule="rul-000004", fn="FUN_IS_TERMINATE"),
        R("Tính đúng đắn", 0.20, ["O_C_CORRECTNESS"],
          "Thiếu đặc trưng đúng đắn với mọi input hợp lệ.",
          "Đúng đắn bị mô tả như chỉ đúng trên một ví dụ.",
          rule="rul-000005", fn="FUN_PRODUCES_CORRECT_OUTPUT"),
        R("Tính hiệu quả", 0.20, ["O_C_EFFICIENCY", "O_C_BIG_O"],
          "Thiếu hiệu quả và đánh giá bằng Big-O / best-average-worst.",
          "Hiệu quả bị nhầm với 'chạy được' mà không đo chi phí.",
          rule="rul-000007", fn="FUN_HAS_COMPLEXITY"),
        R("Tính khả thi", 0.20, ["O_C_FEASIBILITY"],
          "Thiếu đặc trưng khả thi: bước thực hiện được trên máy.",
          "Khả thi bị nhầm với tính đúng đắn.",
          rule="rul-000006", fn="FUN_IS_IMPLEMENTABLE"),
    ],
    "Q_C1_R_PRO_E": [
        R("Địa chỉ cơ sở", 0.30, ["O_C_ARRAY", "O_C_ARRAY_ACCESS"],
          "Thiếu bước xác định địa chỉ cơ sở của mảng.",
          "Không dùng địa chỉ cơ sở khi truy cập A[2]."),
        R("Công thức chỉ số", 0.40, ["O_C_ARRAY_ACCESS", "O_C_1D_ARRAY"],
          "Thiếu công thức địa chỉ = cơ sở + index × kích thước phần tử.",
          "Tính địa chỉ bằng cách duyệt tuần tự.",
          fn="FUN_ACCESS_ARRAY"),
        R("Đọc O(1)", 0.30, ["O_C_CONSTANT_TIME"],
          "Thiếu kết quả A[2] = 30 hoặc độ phức tạp O(1).",
          "Kết luận sai giá trị hoặc cho rằng truy cập là O(n).",
          rule="rul-000034"),
    ],
    "Q_C1_R_PRO_M": [
        R("Duyệt theo chỉ số", 0.40, ["O_C_ARRAY_TRAVERSE", "O_C_1D_ARRAY"],
          "Thiếu lần lượt đọc từng phần tử theo chỉ số tăng dần.",
          "Bỏ sót phần tử hoặc duyệt không theo thứ tự chỉ số.",
          fn="FUN_TRAVERSE_ARRAY"),
        R("Điều kiện dừng", 0.30, ["O_C_ARRAY"],
          "Thiếu điều kiện dừng khi i đạt n.",
          "Dừng sai (quá sớm hoặc vượt biên).",
          fn="FUN_LENGTH"),
        R("Độ phức tạp tuyến tính", 0.30, ["O_C_LINEAR_TIME"],
          "Thiếu kết luận duyệt mảng có độ phức tạp O(n).",
          "Gán O(1) hoặc O(n^2) cho duyệt toàn mảng."),
    ],
    "Q_C1_R_PRO_H": [
        R("Dịch phải phần tử", 0.40, ["O_C_ARRAY_INSERT"],
          "Thiếu bước dịch các phần tử bên phải vị trí chèn.",
          "Dịch trái hoặc chèn đè mà không dời phần tử.",
          fn="FUN_INSERT_ARRAY"),
        R("Ghi giá trị mới", 0.30, ["O_C_ARRAY_INSERT", "O_C_ARRAY"],
          "Thiếu bước ghi giá trị mới vào vị trí chèn.",
          "Ghi sai chỉ số chèn."),
        R("Cập nhật kích thước", 0.30, ["O_C_ARRAY"],
          "Thiếu cập nhật n sau khi chèn.",
          "Không tăng n hoặc tăng sai."),
    ],
    "Q_C1_R_APP_E": [
        R("Khai báo hàm duyệt", 0.30, ["O_C_ARRAY_TRAVERSE", "O_C_PSEUDOCODE"],
          "Thiếu hàm traverse(A, n).",
          "Sai tham số hoặc không viết dạng hàm.",
          fn="FUN_TRAVERSE_ARRAY"),
        R("Vòng lặp 0..n-1", 0.40, ["O_C_ARRAY", "O_C_ARRAY_TRAVERSE"],
          "Thiếu vòng FOR i = 0 TO n-1.",
          "Sai cận vòng lặp (1..n hoặc n inclusive).",
          fn="FUN_LENGTH"),
        R("Xuất từng phần tử", 0.30, ["O_C_ARRAY_ACCESS"],
          "Thiếu việc xuất A[i] trong vòng lặp.",
          "Truy cập sai chỉ số hoặc không duyệt hết."),
    ],
    "Q_C1_R_APP_M": [
        R("Khởi tạo max", 0.30, ["O_C_ARRAY_TRAVERSE", "O_C_ARRAY"],
          "Thiếu khởi tạo max từ A[0].",
          "Khởi tạo max sai (0 khi có số âm, hoặc không khởi tạo)."),
        R("Duyệt so sánh", 0.40, ["O_C_ARRAY_TRAVERSE"],
          "Thiếu vòng duyệt và cập nhật khi A[i] lớn hơn max.",
          "So sánh ngược hoặc không cập nhật max.",
          fn="FUN_TRAVERSE_ARRAY"),
        R("Trả về kết quả", 0.30, ["O_C_OUTPUT"],
          "Thiếu RETURN maxValue.",
          "Trả về chỉ số thay vì giá trị lớn nhất."),
    ],
    "Q_C1_R_APP_H": [
        R("Mảng rỗng", 0.30, ["O_C_ARRAY"],
          "Thiếu xử lý n = 0 (trả về 0).",
          "Không xét mảng rỗng.",
          rule="rul-000013", fn="FUN_IS_EMPTY"),
        R("Cộng dồn khi duyệt", 0.40, ["O_C_ARRAY_TRAVERSE"],
          "Thiếu vòng lặp cộng A[i] vào tổng.",
          "Sai cận vòng lặp hoặc không cộng đúng phần tử.",
          fn="FUN_TRAVERSE_ARRAY"),
        R("Trả về tổng", 0.30, ["O_C_OUTPUT"],
          "Thiếu RETURN tổng.",
          "Trả về trung bình hoặc giá trị khác tổng."),
    ],
    "Q_C1_U_DES_E": [
        R("Cấu trúc tĩnh", 0.35, ["O_C_STATIC_STRUCTURE", "O_C_ARRAY"],
          "Thiếu ý cấu trúc tĩnh có kích thước cố định, điển hình là mảng.",
          "Gán tính tĩnh cho danh sách liên kết.",
          rule="rul-000023", fn="FUN_HAS_FIXED_SIZE"),
        R("Cấu trúc động", 0.35, ["O_C_DYNAMIC_STRUCTURE", "O_C_LINKED_LIST"],
          "Thiếu ý cấu trúc động tăng/giảm khi chạy, điển hình là DSLK.",
          "Nhầm động với tĩnh.",
          rule="rul-000024", fn="FUN_CAN_GROW"),
        R("So sánh chi phí", 0.30, ["O_C_ARRAY", "O_C_LINKED_LIST"],
          "Thiếu so sánh truy cập O(1) vs O(n) và chèn/xóa.",
          "So sánh độ phức tạp ngược giữa mảng và DSLK.",
          fn="FUN_COMPARE_ARRAY_AND_LIST"),
    ],
    "Q_C1_U_DES_M": [
        R("Độ phức tạp thời gian", 0.40, ["O_C_TIME_COMPLEXITY", "O_C_EXECUTION_TIME"],
          "Thiếu đo thời gian theo số phép toán theo n (best/average/worst).",
          "Nhầm thời gian với thời gian đồng hồ tuyệt đối."),
        R("Độ phức tạp không gian", 0.35, ["O_C_SPACE_COMPLEXITY", "O_C_MEMORY_USAGE"],
          "Thiếu đo bộ nhớ phụ ngoài input.",
          "Nhầm không gian với thời gian."),
        R("Ký hiệu Big-O", 0.25, ["O_C_BIG_O", "O_C_COMPLEXITY"],
          "Thiếu dùng Big-O để biểu diễn hai loại phức tạp.",
          "Dùng Big-O không đúng vai trò cận trên.",
          fn="FUN_HAS_COMPLEXITY"),
    ],
    "Q_C1_U_DES_H": [
        R("Truy cập trực tiếp O(1)", 0.40, ["O_C_DIRECT_ACCESS_DATA_STRUCTURE", "O_C_ARRAY"],
          "Thiếu ý truy cập trực tiếp tính được địa chỉ, O(1).",
          "Gán truy cập trực tiếp cho DSLK.",
          rule="rul-000034", fn="FUN_IS_DIRECT_ACCESS"),
        R("Truy cập tuần tự O(n)", 0.35, ["O_C_SEQUENTIAL_ACCESS_DATA_STRUCTURE", "O_C_LINKED_LIST"],
          "Thiếu ý truy cập tuần tự phải đi từ đầu, O(n).",
          "Cho rằng DSLK truy cập ngẫu nhiên O(1)."),
        R("Công thức địa chỉ mảng", 0.25, ["O_C_ARRAY_ACCESS"],
          "Thiếu A[i] = cơ sở + i × kích thước.",
          "Mô tả sai cách tính địa chỉ.",
          fn="FUN_ACCESS_ARRAY"),
    ],
    "Q_C1_U_PRO_E": [
        R("Không duyệt phần tử trước", 0.40, ["O_C_ARRAY_ACCESS"],
          "Thiếu giải thích không cần duyệt A[0], A[1] để lấy A[2].",
          "Giải thích như duyệt tuần tự.",
          fn="FUN_ACCESS_ARRAY"),
        R("Phép tính địa chỉ", 0.35, ["O_C_CONSTANT_TIME"],
          "Thiếu ý chỉ số nhân kích thước rồi cộng cơ sở là O(1).",
          "Cho rằng số phép toán tăng theo n.",
          rule="rul-000034"),
        R("Khác DSLK", 0.25, ["O_C_LINKED_LIST", "O_C_HEAD_POINTER"],
          "Thiếu so sánh với DSLK phải đi từ head.",
          "So sánh sai với DSLK."),
    ],
    "Q_C1_U_PRO_M": [
        R("Hai vòng lặp", 0.35, ["O_C_THEORETICAL_ANALYSIS", "O_C_2D_ARRAY"],
          "Thiếu nhận hai vòng i, j mỗi vòng n lần.",
          "Đếm sai số vòng lặp."),
        R("Tổng n^2 so sánh", 0.40, ["O_C_QUADRATIC_TIME", "O_C_TIME_COMPLEXITY"],
          "Thiếu tổng số so sánh = n^2 và Big-O O(n^2).",
          "Kết luận O(n) hoặc O(n log n)."),
        R("Phân tích lý thuyết", 0.25, ["O_C_THEORETICAL_ANALYSIS", "O_C_EXECUTION_TIME"],
          "Thiếu ý đây là phân tích lý thuyết theo số phép toán, không đo đồng hồ.",
          "Nhầm với phân tích thực nghiệm."),
    ],
    "Q_C1_U_PRO_H": [
        R("Dịch từ cuối", 0.40, ["O_C_ARRAY_INSERT"],
          "Thiếu dịch phần tử từ cuối về vị trí 0.",
          "Dịch sai hướng khi chèn đầu mảng.",
          fn="FUN_INSERT_ARRAY"),
        R("Chèn tại 0", 0.35, ["O_C_ARRAY"],
          "Thiếu gán giá trị mới vào A[0].",
          "Chèn vào chỉ số khác 0."),
        R("Chi phí O(n)", 0.25, ["O_C_LINEAR_TIME"],
          "Thiếu kết luận chèn đầu mảng là O(n) vì dịch n phần tử.",
          "Kết luận O(1) như chèn đầu DSLK."),
    ],
    "Q_C1_U_APP_E": [
        R("Kiểm tra index < 0", 0.30, ["O_C_ARRAY_ACCESS"],
          "Thiếu kiểm tra index < 0 và trả về -1.",
          "Điều kiện cận dưới sai.",
          rule="rul-000010"),
        R("Kiểm tra index >= n", 0.30, ["O_C_ARRAY"],
          "Thiếu kiểm tra index >= n và trả về -1.",
          "Dùng > n hoặc không so với n.",
          rule="rul-000010", fn="FUN_LENGTH"),
        R("Trả về A[index]", 0.40, ["O_C_ARRAY_ACCESS"],
          "Thiếu RETURN A[index] khi chỉ số hợp lệ.",
          "Duyệt tuần tự thay vì truy cập trực tiếp.",
          fn="FUN_ACCESS_ARRAY"),
    ],
    "Q_C1_U_APP_M": [
        R("Dịch phải từ n", 0.40, ["O_C_ARRAY_INSERT"],
          "Thiếu vòng dịch A[i] = A[i-1] từ cuối về index.",
          "Dịch trái hoặc sai hướng vòng lặp.",
          fn="FUN_INSERT_ARRAY"),
        R("Ghi giá trị và tăng n", 0.35, ["O_C_ARRAY_INSERT"],
          "Thiếu A[index] = value và n = n + 1.",
          "Quên tăng n hoặc ghi sai chỗ."),
        R("Trả về n mới", 0.25, ["O_C_OUTPUT"],
          "Thiếu RETURN n mới.",
          "Không trả về kích thước sau chèn."),
    ],
    "Q_C1_U_APP_H": [
        R("Hai vòng hàng cột", 0.45, ["O_C_2D_ARRAY", "O_C_ARRAY_TRAVERSE"],
          "Thiếu hai vòng FOR theo hàng rồi cột.",
          "Chỉ một vòng hoặc đảo hàng/cột sai.",
          fn="FUN_TRAVERSE_ARRAY"),
        R("Xuất A[i][j]", 0.30, ["O_C_ARRAY_ACCESS", "O_C_2D_ARRAY"],
          "Thiếu xuất từng A[i][j].",
          "Truy cập sai chỉ số 2 chiều."),
        R("Cận rows và cols", 0.25, ["O_C_2D_ARRAY"],
          "Thiếu cận 0..rows-1 và 0..cols-1.",
          "Sai cận vòng lặp."),
    ],
    "Q_C1_AP_DES_E": [
        R("Chọn mảng", 0.40, ["O_C_ARRAY", "O_C_DIRECT_ACCESS_DATA_STRUCTURE"],
          "Thiếu kết luận chọn mảng vì cần truy cập ngẫu nhiên.",
          "Chọn DSLK cho nhu cầu truy cập theo chỉ số.",
          rule="rul-000035", fn="FUN_RECOMMEND_DATA_STRUCTURE"),
        R("Truy cập O(1)", 0.35, ["O_C_ARRAY_ACCESS", "O_C_CONSTANT_TIME"],
          "Thiếu lập luận A[i] là O(1).",
          "Cho rằng mảng truy cập O(n).",
          rule="rul-000034", fn="FUN_IS_DIRECT_ACCESS"),
        R("DSLK truy cập O(n)", 0.25, ["O_C_LINKED_LIST", "O_C_LINEAR_TIME"],
          "Thiếu so sánh DSLK phải duyệt từ head, O(n).",
          "Nêu DSLK truy cập ngẫu nhiên nhanh hơn mảng."),
    ],
    "Q_C1_AP_DES_M": [
        R("Dữ liệu là mảng", 0.35, ["O_C_DATA", "O_C_ARRAY"],
          "Thiếu xác định dữ liệu là mảng điểm.",
          "Không nêu cấu trúc dữ liệu của chương trình.",
          rule="rul-000008", fn="FUN_INCLUDES_DATA"),
        R("Thuật toán duyệt tính TB", 0.35, ["O_C_ALGORITHM", "O_C_ARRAY_TRAVERSE"],
          "Thiếu thuật toán duyệt, cộng dồn, chia n.",
          "Nêu dữ liệu nhưng không nêu thuật toán.",
          fn="FUN_INCLUDES_ALGORITHM"),
        R("Chương trình kết hợp", 0.30, ["O_C_PROGRAM"],
          "Thiếu ý chương trình = thuật toán + dữ liệu.",
          "Tách chương trình khỏi thuật toán hoặc dữ liệu.",
          rule="rul-000008"),
    ],
    "Q_C1_AP_DES_H": [
        R("Vét cạn O(n)", 0.40, ["O_C_BRUTE_FORCE", "O_C_LINEAR_SEARCH"],
          "Thiếu chọn vét cạn/tuyến tính khi dữ liệu chưa có thứ tự.",
          "Dùng nhị phân trên dữ liệu chưa sắp.",
          rule="rul-000026", fn="FUN_RECOMMEND_SEARCH"),
        R("Chia để trị O(log n)", 0.40, ["O_C_DIVIDE_AND_CONQUER", "O_C_BINARY_SEARCH"],
          "Thiếu chia để trị / tìm kiếm nhị phân khi đã sắp.",
          "Không nêu chia đôi khoảng tìm kiếm.",
          rule="rul-000025"),
        R("Kỹ thuật thiết kế", 0.20, ["O_C_ALGORITHM_DESIGN"],
          "Thiếu gắn mỗi bài toán với kỹ thuật thiết kế tương ứng.",
          "Nhầm brute force với divide and conquer."),
    ],
    "Q_C1_AP_PRO_E": [
        R("Dịch phải từ cuối", 0.40, ["O_C_ARRAY_INSERT"],
          "Thiếu các bước dịch A[3]→A[4], A[2]→A[3], A[1]→A[2].",
          "Dịch sai hướng hoặc thiếu bước.",
          fn="FUN_INSERT_ARRAY"),
        R("Ghi 8 vào index 1", 0.35, ["O_C_ARRAY_INSERT"],
          "Thiếu A[1] = 8 và n = 5.",
          "Ghi sai vị trí hoặc không tăng n."),
        R("Kết quả mảng", 0.25, ["O_C_ARRAY"],
          "Thiếu kết quả [1, 8, 3, 4, 5].",
          "Mảng kết quả sai."),
    ],
    "Q_C1_AP_PRO_M": [
        R("Xác định phần tử xóa", 0.25, ["O_C_ARRAY_DELETE"],
          "Thiếu xác định phần tử bị xóa A[2].",
          "Xóa nhầm chỉ số."),
        R("Dịch trái", 0.45, ["O_C_ARRAY_DELETE"],
          "Thiếu dịch trái các phần tử bên phải vị trí xóa.",
          "Dịch phải hoặc không dồn mảng.",
          fn="FUN_DELETE_ARRAY"),
        R("Giảm n", 0.30, ["O_C_ARRAY"],
          "Thiếu n = n - 1 và dãy còn 4 phần tử.",
          "Không giảm n sau khi xóa."),
    ],
    "Q_C1_AP_PRO_H": [
        R("Truy cập A[1][0]", 0.35, ["O_C_2D_ARRAY", "O_C_ARRAY_ACCESS"],
          "Thiếu truy cập đúng hàng 1 cột 0.",
          "Nhầm hàng/cột."),
        R("Cập nhật giá trị", 0.40, ["O_C_ARRAY_UPDATE", "O_C_2D_ARRAY"],
          "Thiếu gán A[1][0] = 9.",
          "Cập nhật ô khác.",
          fn="FUN_UPDATE_ARRAY"),
        R("Duyệt lại ma trận", 0.25, ["O_C_ARRAY_TRAVERSE"],
          "Thiếu duyệt lại các phần tử sau cập nhật.",
          "Không kiểm tra ma trận sau khi gán."),
    ],
    "Q_C1_AP_APP_E": [
        R("Dịch trái từ index", 0.45, ["O_C_ARRAY_DELETE"],
          "Thiếu FOR i = index TO n-2: A[i] = A[i+1].",
          "Sai hướng dịch hoặc sai cận.",
          fn="FUN_DELETE_ARRAY"),
        R("Giảm n và trả về", 0.35, ["O_C_ARRAY"],
          "Thiếu n = n - 1 và RETURN n.",
          "Không trả về kích thước mới."),
        R("Chỉ số hợp lệ ngầm", 0.20, ["O_C_ARRAY_ACCESS"],
          "Thiếu giả định 0 ≤ index < n như đề.",
          "Xóa ngoài biên mà không nêu hệ quả.",
          rule="rul-000010"),
    ],
    "Q_C1_AP_APP_M": [
        R("Kiểm tra biên index", 0.40, ["O_C_ARRAY_UPDATE"],
          "Thiếu IF index < 0 hoặc index >= n rồi RETURN 0.",
          "Không kiểm tra biên trước khi ghi.",
          rule="rul-000010", fn="FUN_LENGTH"),
        R("Gán A[index]", 0.40, ["O_C_ARRAY_UPDATE", "O_C_ARRAY_ACCESS"],
          "Thiếu A[index] = value.",
          "Cập nhật sai phần tử.",
          fn="FUN_UPDATE_ARRAY"),
        R("Trả về thành công", 0.20, ["O_C_OUTPUT"],
          "Thiếu RETURN 1 khi cập nhật thành công.",
          "Không phân biệt thành công/thất bại."),
    ],
    "Q_C1_AP_APP_H": [
        R("Hai vòng duyệt 2D", 0.40, ["O_C_2D_ARRAY", "O_C_ARRAY_TRAVERSE"],
          "Thiếu hai vòng FOR theo hàng và cột.",
          "Chỉ duyệt một chiều."),
        R("So sánh với khóa", 0.35, ["O_C_LINEAR_SEARCH", "O_C_2D_ARRAY"],
          "Thiếu so sánh A[i][j] = x và RETURN 1 khi thấy.",
          "Không dừng khi tìm thấy hoặc so sánh sai.",
          fn="FUN_LINEAR_SEARCH"),
        R("Không thấy thì 0", 0.25, ["O_C_OUTPUT"],
          "Thiếu RETURN 0 khi duyệt hết không có x.",
          "Không có nhánh thất bại."),
    ],
    "Q_C2_R_DES_E": [
        R("Duyệt tuần tự", 0.40, ["O_C_LINEAR_SEARCH", "O_C_SEARCHING_ALGORITHM"],
          "Thiếu ý duyệt tuần tự từng phần tử so với khóa.",
          "Mô tả như chia đôi khoảng.",
          fn="FUN_LINEAR_SEARCH"),
        R("Không cần sắp", 0.30, ["O_C_LINEAR_SEARCH", "O_C_ARRAY"],
          "Thiếu ý áp dụng được cho dữ liệu chưa sắp.",
          "Yêu cầu mảng đã sắp cho tuyến tính.",
          rule="rul-000012"),
        R("Độ phức tạp O(1)/O(n)", 0.30, ["O_C_CONSTANT_TIME", "O_C_LINEAR_TIME"],
          "Thiếu best O(1) và average/worst O(n), không gian O(1).",
          "Gán O(log n) cho tìm kiếm tuyến tính."),
    ],
    "Q_C2_R_DES_M": [
        R("Mảng đã sắp", 0.35, ["O_C_BINARY_SEARCH", "O_C_1D_ARRAY"],
          "Thiếu tiền điều kiện mảng một chiều đã sắp.",
          "Áp dụng nhị phân cho dữ liệu chưa sắp.",
          rule="rul-000011", fn="FUN_IS_SORTED"),
        R("Thu hẹp nửa", 0.35, ["O_C_BINARY_SEARCH", "O_C_DIVIDE_AND_CONQUER"],
          "Thiếu so sánh với phần tử giữa rồi loại một nửa.",
          "Duyệt tuần tự thay vì chia đôi.",
          fn="FUN_BINARY_SEARCH"),
        R("O(1) và O(log n)", 0.30, ["O_C_CONSTANT_TIME", "O_C_LOGARITHMIC_TIME"],
          "Thiếu best O(1), average/worst O(log n), không gian O(1).",
          "Gán O(n) cho worst-case nhị phân."),
    ],
    "Q_C2_R_DES_H": [
        R("Đổi chỗ cặp kề", 0.40, ["O_C_BUBBLE_SORT", "O_C_SORTING_ALGORITHM"],
          "Thiếu ý so sánh cặp liền kề và đổi chỗ nếu sai thứ tự.",
          "Nhầm với chọn min như Selection Sort.",
          fn="FUN_BUBBLE_SORT"),
        R("Best O(n) có cờ", 0.30, ["O_C_LINEAR_TIME", "O_C_BUBBLE_SORT"],
          "Thiếu best-case O(n) khi có cờ và dữ liệu đã sắp.",
          "Nêu Bubble luôn O(n^2) kể cả best có cờ.",
          rule="rul-000037"),
        R("Worst O(n^2)", 0.30, ["O_C_QUADRATIC_TIME"],
          "Thiếu average/worst O(n^2) và không gian O(1).",
          "Gán O(n log n) cho Bubble Sort."),
    ],
    "Q_C2_R_PRO_E": [
        R("So sánh lần lượt", 0.45, ["O_C_LINEAR_SEARCH"],
          "Thiếu các bước i=0, i=1 so sánh với 7.",
          "Bỏ bước hoặc nhảy có điều kiện như nhị phân.",
          fn="FUN_LINEAR_SEARCH"),
        R("Dừng khi thấy", 0.35, ["O_C_LINEAR_SEARCH"],
          "Thiếu dừng khi A[1] = 7.",
          "Duyệt hết mảng dù đã thấy khóa."),
        R("Kết quả chỉ số 1", 0.20, ["O_C_OUTPUT"],
          "Thiếu kết quả tìm thấy tại chỉ số 1.",
          "Trả sai chỉ số."),
    ],
    "Q_C2_R_PRO_M": [
        R("Khởi tạo left right mid", 0.30, ["O_C_BINARY_SEARCH"],
          "Thiếu left, right, mid ban đầu.",
          "Không duy trì khoảng tìm kiếm.",
          fn="FUN_BINARY_SEARCH"),
        R("Thu hẹp theo so sánh", 0.45, ["O_C_BINARY_SEARCH"],
          "Thiếu cập nhật right hoặc left sau mỗi lần so với A[mid].",
          "Thu hẹp sai nửa.",
          rule="rul-000011"),
        R("Tìm thấy tại mid", 0.25, ["O_C_OUTPUT"],
          "Thiếu kết luận tìm thấy khi A[mid] bằng khóa.",
          "Kết luận sai chỉ số."),
    ],
    "Q_C2_R_PRO_H": [
        R("So sánh cặp kề", 0.40, ["O_C_BUBBLE_SORT"],
          "Thiếu các bước so sánh 5-1, 5-4, 5-2.",
          "Không so sánh cặp liền kề.",
          fn="FUN_BUBBLE_SORT"),
        R("Đổi chỗ", 0.35, ["O_C_BUBBLE_SORT"],
          "Thiếu đổi chỗ khi phần tử trái lớn hơn phải.",
          "Không đổi chỗ hoặc đổi sai.",
          fn="FUN_SWAP"),
        R("Phần tử lớn về cuối", 0.25, ["O_C_BUBBLE_SORT"],
          "Thiếu kết quả sau một lượt: phần tử lớn nhất ở cuối.",
          "Không nêu phần tử được 'nổi' về cuối."),
    ],
    "Q_C2_R_APP_E": [
        R("Vòng duyệt n phần tử", 0.35, ["O_C_LINEAR_SEARCH", "O_C_ARRAY"],
          "Thiếu FOR i = 0 TO n-1.",
          "Sai cận vòng lặp.",
          fn="FUN_LINEAR_SEARCH"),
        R("RETURN khi thấy", 0.40, ["O_C_LINEAR_SEARCH"],
          "Thiếu IF A[i] = x THEN RETURN i.",
          "Không trả về chỉ số khi tìm thấy."),
        R("Không thấy trả -1", 0.25, ["O_C_OUTPUT"],
          "Thiếu RETURN -1 khi hết mảng.",
          "Không có nhánh thất bại."),
    ],
    "Q_C2_R_APP_M": [
        R("Tiền điều kiện đã sắp", 0.25, ["O_C_BINARY_SEARCH"],
          "Thiếu giả định/vòng WHILE trên khoảng left..right.",
          "Viết như tìm tuyến tính.",
          rule="rul-000011", fn="FUN_IS_SORTED"),
        R("Tính mid và so sánh", 0.45, ["O_C_BINARY_SEARCH"],
          "Thiếu mid = (left+right)/2 và ba nhánh = / < / >.",
          "Cập nhật left/right sai.",
          fn="FUN_BINARY_SEARCH"),
        R("Không thấy trả -1", 0.30, ["O_C_OUTPUT"],
          "Thiếu RETURN -1 khi left > right.",
          "Lặp vô hạn hoặc không trả thất bại."),
    ],
    "Q_C2_R_APP_H": [
        R("Hai vòng i và j", 0.35, ["O_C_BUBBLE_SORT"],
          "Thiếu hai vòng FOR i, j cho các lượt nổi bọt.",
          "Thiếu vòng trong so cặp kề.",
          fn="FUN_BUBBLE_SORT"),
        R("Đổi chỗ nếu nghịch thế", 0.40, ["O_C_BUBBLE_SORT"],
          "Thiếu IF A[j] > A[j+1] rồi đổi chỗ.",
          "So sánh ngược chiều sắp tăng.",
          fn="FUN_SWAP"),
        R("Cờ swapped", 0.25, ["O_C_BUBBLE_SORT", "O_C_LINEAR_TIME"],
          "Thiếu cờ swapped để best-case O(n).",
          "Không có cơ chế dừng sớm.",
          rule="rul-000037"),
    ],
    "Q_C2_U_DES_E": [
        R("Cần mảng đã sắp", 0.45, ["O_C_BINARY_SEARCH", "O_C_1D_ARRAY"],
          "Thiếu giải thích nhị phân dựa trên thứ tự để loại một nửa.",
          "Cho rằng nhị phân luôn đúng trên mảng chưa sắp.",
          rule="rul-000011", fn="FUN_IS_SORTED"),
        R("Loại nhầm khóa", 0.35, ["O_C_BINARY_SEARCH"],
          "Thiếu ý nếu chưa sắp, loại nửa có thể loại luôn khóa.",
          "Không nêu rủi ro kết luận 'không có' sai."),
        R("Phải dùng tuyến tính", 0.20, ["O_C_LINEAR_SEARCH"],
          "Thiếu kết luận khi chưa sắp phải dùng tuyến tính.",
          "Vẫn khuyên nhị phân cho dữ liệu chưa sắp.",
          rule="rul-000026", fn="FUN_RECOMMEND_SEARCH"),
    ],
    "Q_C2_U_DES_M": [
        R("Ổn định", 0.40, ["O_C_SORTING_ALGORITHM"],
          "Thiếu định nghĩa ổn định: giữ thứ tự tương đối khóa bằng nhau.",
          "Nhầm ổn định với tại chỗ.",
          fn="FUN_IS_STABLE"),
        R("Tại chỗ", 0.35, ["O_C_SPACE_COMPLEXITY"],
          "Thiếu định nghĩa tại chỗ: bộ nhớ phụ O(1).",
          "Nhầm tại chỗ với thời gian O(1).",
          fn="FUN_IS_IN_PLACE"),
        R("Bubble Insertion Selection", 0.25, ["O_C_BUBBLE_SORT", "O_C_INSERTION_SORT", "O_C_SELECTION_SORT"],
          "Thiếu phân loại ổn định/tại chỗ của ba thuật toán cơ bản.",
          "Gán Selection Sort là ổn định.",
          rule="rul-000038"),
    ],
    "Q_C2_U_DES_H": [
        R("Selection luôn O(n^2)", 0.35, ["O_C_SELECTION_SORT", "O_C_QUADRATIC_TIME"],
          "Thiếu ý Selection luôn quét đoạn chưa sắp, kể cả dãy đã tăng.",
          "Nêu Selection có best O(n).",
          fn="FUN_SELECTION_SORT"),
        R("Bubble best O(n)", 0.35, ["O_C_BUBBLE_SORT", "O_C_LINEAR_TIME"],
          "Thiếu Bubble có cờ: best O(n) khi đã sắp.",
          "Không phân biệt best của Bubble và Selection.",
          rule="rul-000037"),
        R("Insertion best O(n)", 0.30, ["O_C_INSERTION_SORT", "O_C_LINEAR_TIME"],
          "Thiếu Insertion gần sắp thì gần O(n).",
          "Cho Insertion luôn O(n^2) không có best tuyến tính.",
          rule="rul-000027"),
    ],
    "Q_C2_U_PRO_E": [
        R("Trúng mid ngay", 0.40, ["O_C_BINARY_SEARCH", "O_C_CONSTANT_TIME"],
          "Thiếu bước A[mid] = x nên dừng, đây là best-case O(1).",
          "Vẫn chia thêm nửa dù đã thấy khóa.",
          fn="FUN_BINARY_SEARCH"),
        R("Loại nửa còn lại", 0.35, ["O_C_BINARY_SEARCH"],
          "Thiếu giải thích nếu A[mid] > x thì loại nửa phải.",
          "Thu hẹp sai hướng."),
        R("Tiền điều kiện sắp", 0.25, ["O_C_1D_ARRAY"],
          "Thiếu nhấn mạnh tính đúng vì mảng đã sắp.",
          "Bỏ qua tiền điều kiện.",
          rule="rul-000011", fn="FUN_IS_SORTED"),
    ],
    "Q_C2_U_PRO_M": [
        R("Chọn min mỗi lượt", 0.45, ["O_C_SELECTION_SORT"],
          "Thiếu mỗi lượt tìm min trên đoạn chưa sắp rồi đổi chỗ.",
          "Mô tả như nổi bọt cặp kề.",
          fn="FUN_SELECTION_SORT"),
        R("Đoạn đã sắp tăng dần", 0.35, ["O_C_SELECTION_SORT"],
          "Thiếu ý đầu mảng dần trở thành đoạn đã sắp.",
          "Không nêu bất biến đoạn đã chọn."),
        R("Worst không cải thiện", 0.20, ["O_C_QUADRATIC_TIME"],
          "Thiếu nhận số so sánh vẫn bậc hai.",
          "Nêu Selection có best O(n)."),
    ],
    "Q_C2_U_PRO_H": [
        R("Chèn key vào đoạn sắp", 0.45, ["O_C_INSERTION_SORT"],
          "Thiếu lấy key và dịch phần tử lớn hơn sang phải rồi chèn.",
          "Nhầm với đổi min như Selection.",
          fn="FUN_INSERTION_SORT"),
        R("Đoạn trái đã sắp", 0.30, ["O_C_INSERTION_SORT"],
          "Thiếu bất biến: đoạn trái sau mỗi bước đã sắp.",
          "Không nêu đoạn đã sắp tăng dần."),
        R("Dãy ngược gần worst", 0.25, ["O_C_QUADRATIC_TIME"],
          "Thiếu nhận dãy gần giảm là gần worst-case.",
          "Nêu đây là best-case."),
    ],
    "Q_C2_U_APP_E": [
        R("Mảng rỗng", 0.25, ["O_C_ARRAY"],
          "Thiếu IF n = 0 RETURN -1.",
          "Không xử lý mảng rỗng.",
          rule="rul-000013", fn="FUN_IS_EMPTY"),
        R("Duyệt tuyến tính", 0.45, ["O_C_LINEAR_SEARCH"],
          "Thiếu FOR so sánh A[i] = x rồi RETURN i.",
          "Dùng nhị phân hoặc sai vòng lặp.",
          fn="FUN_LINEAR_SEARCH"),
        R("Không thấy -1", 0.30, ["O_C_OUTPUT"],
          "Thiếu RETURN -1 khi hết mảng.",
          "Không có nhánh thất bại."),
    ],
    "Q_C2_U_APP_M": [
        R("Vòng i từ 1", 0.25, ["O_C_INSERTION_SORT"],
          "Thiếu FOR i = 1 TO n-1 lấy key = A[i].",
          "Bắt đầu i = 0 hoặc không lấy key.",
          fn="FUN_INSERTION_SORT"),
        R("Dịch khi A[j] > key", 0.45, ["O_C_INSERTION_SORT"],
          "Thiếu WHILE dịch A[j] sang phải khi lớn hơn key.",
          "Dịch sai điều kiện hoặc không giảm j."),
        R("Chèn key", 0.30, ["O_C_INSERTION_SORT"],
          "Thiếu A[j+1] = key sau vòng dịch.",
          "Mất key hoặc chèn sai chỗ."),
    ],
    "Q_C2_U_APP_H": [
        R("Lượt i chọn minIndex", 0.35, ["O_C_SELECTION_SORT"],
          "Thiếu FOR i và khởi tạo minIndex = i.",
          "Không giữ vị trí min.",
          fn="FUN_SELECTION_SORT"),
        R("Quét j tìm min", 0.40, ["O_C_SELECTION_SORT"],
          "Thiếu vòng j = i+1..n-1 cập nhật minIndex.",
          "Quét sai đoạn chưa sắp."),
        R("Đổi chỗ một lần", 0.25, ["O_C_SELECTION_SORT"],
          "Thiếu đổi A[i] với A[minIndex] mỗi lượt.",
          "Đổi trong vòng j như Bubble.",
          fn="FUN_SWAP"),
    ],
    "Q_C2_AP_DES_E": [
        R("Chọn tuyến tính", 0.45, ["O_C_LINEAR_SEARCH", "O_C_SEARCHING_ALGORITHM"],
          "Thiếu chọn Linear Search vì A chưa sắp.",
          "Chọn Binary Search trên mảng chưa sắp.",
          rule="rul-000026", fn="FUN_RECOMMEND_SEARCH"),
        R("Nhị phân cần đã sắp", 0.30, ["O_C_BINARY_SEARCH"],
          "Thiếu lập luận Binary Search yêu cầu mảng đã sắp.",
          "Bỏ qua tiền điều kiện nhị phân.",
          rule="rul-000011", fn="FUN_IS_SORTED"),
        R("Kết quả chỉ số 3", 0.25, ["O_C_LINEAR_SEARCH"],
          "Thiếu nêu khóa 3 ở chỉ số 3 sau khi duyệt 2, 5, 4.",
          "Trả sai vị trí."),
    ],
    "Q_C2_AP_DES_M": [
        R("Chọn Insertion n nhỏ", 0.40, ["O_C_INSERTION_SORT"],
          "Thiếu chọn Insertion Sort vì n nhỏ và gần sắp.",
          "Chọn Selection dù gần sắp vẫn O(n^2).",
          rule="rul-000027", fn="FUN_RECOMMEND_SORT"),
        R("Gần best O(n)", 0.30, ["O_C_LINEAR_TIME", "O_C_INSERTION_SORT"],
          "Thiếu ý dữ liệu gần sắp nên gần best-case O(n).",
          "Cho rằng mọi sort cơ bản đều O(n^2) như nhau trên dữ liệu gần sắp."),
        R("Ổn định và tại chỗ", 0.30, ["O_C_INSERTION_SORT"],
          "Thiếu Insertion ổn định và tại chỗ.",
          "Gán không ổn định cho Insertion.",
          fn="FUN_IS_STABLE"),
    ],
    "Q_C2_AP_DES_H": [
        R("Loại Selection không ổn định", 0.45, ["O_C_SELECTION_SORT"],
          "Thiếu loại Selection Sort vì không ổn định.",
          "Chọn Selection khi cần giữ thứ tự cùng điểm.",
          rule="rul-000038", fn="FUN_IS_STABLE"),
        R("Bubble ổn định", 0.25, ["O_C_BUBBLE_SORT"],
          "Thiếu Bubble Sort ổn định nếu chỉ đổi khi >.",
          "Nêu Bubble không ổn định."),
        R("Insertion ổn định", 0.30, ["O_C_INSERTION_SORT"],
          "Thiếu Insertion Sort ổn định, phù hợp yêu cầu.",
          "Loại nhầm Insertion."),
    ],
    "Q_C2_AP_PRO_E": [
        R("Duyệt tới khi thấy", 0.50, ["O_C_LINEAR_SEARCH"],
          "Thiếu các bước i=0..3 so sánh với 3.",
          "Dừng sớm sai hoặc dùng nhị phân.",
          fn="FUN_LINEAR_SEARCH"),
        R("Kết quả chỉ số 3", 0.30, ["O_C_OUTPUT"],
          "Thiếu kết quả chỉ số 3.",
          "Trả sai chỉ số."),
        R("Không xét phần còn lại", 0.20, ["O_C_LINEAR_SEARCH"],
          "Thiếu ý dừng ngay khi thấy, không cần xét phần tử sau.",
          "Duyệt hết mảng dù đã thấy."),
    ],
    "Q_C2_AP_PRO_M": [
        R("Cập nhật khoảng", 0.45, ["O_C_BINARY_SEARCH"],
          "Thiếu các bước cập nhật left/right theo A[mid] so với 3.",
          "Thu hẹp sai nửa.",
          fn="FUN_BINARY_SEARCH"),
        R("Tìm thấy mid = 1", 0.35, ["O_C_BINARY_SEARCH"],
          "Thiếu kết luận A[1] = 3.",
          "Kết luận không tìm thấy."),
        R("Mảng đã sắp", 0.20, ["O_C_1D_ARRAY"],
          "Thiếu dựa trên mảng đã sắp để loại nửa.",
          "Không dùng thứ tự mảng.",
          rule="rul-000011", fn="FUN_IS_SORTED"),
    ],
    "Q_C2_AP_PRO_H": [
        R("Chèn lần lượt", 0.45, ["O_C_INSERTION_SORT"],
          "Thiếu các bước i=1,2,3 với key 2, 4, 1.",
          "Nhầm các bước với Selection/Bubble.",
          fn="FUN_INSERTION_SORT"),
        R("Kết quả tăng dần", 0.30, ["O_C_OUTPUT"],
          "Thiếu mảng cuối [1, 2, 4, 5].",
          "Mảng kết quả chưa sắp."),
        R("Gần worst-case", 0.25, ["O_C_QUADRATIC_TIME"],
          "Thiếu nhận dãy gần giảm nên nhiều lần dịch, gần worst.",
          "Nêu đây là best-case."),
    ],
    "Q_C2_AP_APP_E": [
        R("Không dừng sớm", 0.40, ["O_C_LINEAR_SEARCH", "O_C_ARRAY_TRAVERSE"],
          "Thiếu duyệt hết mảng, không RETURN khi gặp x lần đầu.",
          "Dừng ở lần xuất hiện đầu như linearSearch thường.",
          fn="FUN_LINEAR_SEARCH"),
        R("Đếm khi A[i] = x", 0.40, ["O_C_LINEAR_SEARCH"],
          "Thiếu tăng count khi A[i] = x.",
          "Không đếm hoặc đếm sai."),
        R("RETURN count", 0.20, ["O_C_OUTPUT"],
          "Thiếu RETURN count.",
          "Trả về chỉ số thay vì số lần."),
    ],
    "Q_C2_AP_APP_M": [
        R("Nhánh đã sắp dùng nhị phân", 0.40, ["O_C_BINARY_SEARCH"],
          "Thiếu IF sortedFlag = 1 thì binary search.",
          "Luôn dùng một thuật toán.",
          rule="rul-000025", fn="FUN_BINARY_SEARCH"),
        R("Nhánh chưa sắp dùng tuyến tính", 0.40, ["O_C_LINEAR_SEARCH"],
          "Thiếu nhánh linear search khi chưa sắp.",
          "Dùng nhị phân khi sortedFlag ≠ 1.",
          rule="rul-000026", fn="FUN_LINEAR_SEARCH"),
        R("Cùng kiểu trả về", 0.20, ["O_C_OUTPUT"],
          "Thiếu trả về chỉ số hoặc -1 thống nhất hai nhánh.",
          "Hai nhánh trả kiểu khác nhau."),
    ],
    "Q_C2_AP_APP_H": [
        R("Hàm isSorted", 0.35, ["O_C_ARRAY"],
          "Thiếu hàm kiểm tra dãy tăng (cặp kề).",
          "Không kiểm tra đã sắp trước khi sort.",
          fn="FUN_IS_SORTED"),
        R("Chỉ sort khi chưa sắp", 0.40, ["O_C_BUBBLE_SORT"],
          "Thiếu gọi bubbleSort chỉ khi isSorted trả 0.",
          "Luôn sort kể cả khi đã sắp.",
          fn="FUN_BUBBLE_SORT"),
        R("Best-case O(n) nhờ cờ/kiểm tra", 0.25, ["O_C_LINEAR_TIME", "O_C_BUBBLE_SORT"],
          "Thiếu ý đã sắp thì không cần nhiều lượt đổi chỗ.",
          "Bỏ qua tối ưu best-case.",
          rule="rul-000037"),
    ],
    "Q_C3_R_DES_E": [
        R("Node có data và con trỏ", 0.40, ["O_C_NODE"],
          "Thiếu node gồm data và một/hai con trỏ liên kết.",
          "Mô tả node như ô mảng liên tục.",
          rule="rul-000015", fn="FUN_HAS_DATA"),
        R("Pointer lưu địa chỉ hoặc null", 0.35, ["O_C_POINTER", "O_C_NEXT_POINTER"],
          "Thiếu pointer lưu địa chỉ node hoặc null; next trỏ node kế.",
          "Nhầm pointer với giá trị data.",
          rule="rul-000016", fn="FUN_HAS_NEXT"),
        R("prev trên DSLK đôi", 0.25, ["O_C_PREV_POINTER", "O_C_DOUBLY_LINKED_LIST"],
          "Thiếu prev trỏ node trước trên DSLK đôi.",
          "Gán prev cho DSLK đơn.",
          fn="FUN_HAS_PREV"),
    ],
    "Q_C3_R_DES_M": [
        R("Dãy node next", 0.35, ["O_C_SINGLY_LINKED_LIST", "O_C_NEXT_POINTER"],
          "Thiếu DSLK đơn là dãy node data + next.",
          "Thêm prev vào SLL.",
          rule="rul-000018", fn="FUN_HAS_NEXT"),
        R("Head và node cuối null", 0.35, ["O_C_HEAD_POINTER", "O_C_TAIL_POINTER"],
          "Thiếu head trỏ đầu; node cuối next = null.",
          "Không nêu vai trò head/null terminator.",
          rule="rul-000017", fn="FUN_HAS_HEAD"),
        R("Chèn xóa đầu O(1)", 0.30, ["O_C_SLL_INSERT_HEAD", "O_C_CONSTANT_TIME"],
          "Thiếu chèn/xóa đầu O(1), duyệt một chiều.",
          "Nêu chèn đầu SLL là O(n).",
          rule="rul-000030"),
    ],
    "Q_C3_R_DES_H": [
        R("Node có next và prev", 0.35, ["O_C_DOUBLY_LINKED_LIST", "O_C_NODE"],
          "Thiếu mỗi node có data, next và prev.",
          "Mô tả DLL như SLL.",
          rule="rul-000019", fn="FUN_HAS_PREV"),
        R("Duyệt hai chiều", 0.35, ["O_C_DLL_TRAVERSE", "O_C_TAIL_POINTER"],
          "Thiếu duyệt xuôi từ head hoặc ngược từ tail.",
          "Nêu DLL chỉ duyệt một chiều.",
          rule="rul-000033", fn="FUN_SUPPORTS_BACKWARD"),
        R("Xóa cuối O(1) nếu có tail", 0.30, ["O_C_TAIL_POINTER", "O_C_CONSTANT_TIME"],
          "Thiếu chèn/xóa cuối O(1) nhờ prev của tail.",
          "Cho xóa cuối DLL luôn O(n) như SLL."),
    ],
    "Q_C3_R_PRO_E": [
        R("p = head rồi p.next", 0.45, ["O_C_SLL_TRAVERSE", "O_C_HEAD_POINTER"],
          "Thiếu p xuất phát từ head và tiến p = p.next.",
          "Duyệt theo chỉ số như mảng.",
          fn="FUN_TRAVERSE_LIST"),
        R("Dừng khi p = null", 0.30, ["O_C_POINTER"],
          "Thiếu dừng khi p = null.",
          "Không có điều kiện dừng.",
          fn="FUN_IS_NULL_POINTER"),
        R("Kết quả và O(n)", 0.25, ["O_C_LINEAR_TIME"],
          "Thiếu dãy 3, 5, 8 và độ phức tạp O(n).",
          "Sai dãy duyệt hoặc O(1)."),
    ],
    "Q_C3_R_PRO_M": [
        R("Tạo node mới", 0.25, ["O_C_NODE"],
          "Thiếu tạo node mới với data = 9.",
          "Không tạo node."),
        R("newNode.next = head", 0.40, ["O_C_SLL_INSERT_HEAD", "O_C_NEXT_POINTER"],
          "Thiếu nối newNode.next vào head cũ rồi cập nhật head.",
          "Gán head trước khi nối next, mất danh sách.",
          fn="FUN_INSERT_HEAD"),
        R("O(1) không duyệt", 0.35, ["O_C_CONSTANT_TIME"],
          "Thiếu kết luận O(1), không duyệt danh sách.",
          "Duyệt tới cuối rồi mới chèn đầu.",
          rule="rul-000030"),
    ],
    "Q_C3_R_PRO_H": [
        R("Danh sách rỗng", 0.25, ["O_C_SINGLY_LINKED_LIST"],
          "Thiếu kiểm tra head = null thì dừng.",
          "Xóa head khi danh sách rỗng.",
          rule="rul-000021", fn="FUN_IS_EMPTY"),
        R("head = head.next", 0.45, ["O_C_SLL_DELETE_HEAD", "O_C_HEAD_POINTER"],
          "Thiếu cập nhật head sang node kế trước khi hủy node cũ.",
          "Hủy head mà không giữ next.",
          fn="FUN_DELETE_HEAD"),
        R("O(1)", 0.30, ["O_C_CONSTANT_TIME"],
          "Thiếu độ phức tạp O(1).",
          "Nêu xóa đầu là O(n).",
          rule="rul-000031"),
    ],
    "Q_C3_R_APP_E": [
        R("WHILE p <> null", 0.40, ["O_C_SLL_TRAVERSE"],
          "Thiếu vòng WHILE p khác null.",
          "Dùng FOR theo chỉ số mảng.",
          fn="FUN_TRAVERSE_LIST"),
        R("Xuất data rồi p.next", 0.40, ["O_C_NODE", "O_C_NEXT_POINTER"],
          "Thiếu OUTPUT p.data và p = p.next.",
          "Không tiến con trỏ, lặp vô hạn.",
          fn="FUN_HAS_DATA"),
        R("Bắt đầu từ head", 0.20, ["O_C_HEAD_POINTER"],
          "Thiếu p = head.",
          "Bắt đầu từ node giữa hoặc tail.",
          fn="FUN_HAS_HEAD"),
    ],
    "Q_C3_R_APP_M": [
        R("Gán data", 0.25, ["O_C_NODE"],
          "Thiếu newNode.data = value.",
          "Không gán dữ liệu node mới.",
          fn="FUN_HAS_DATA"),
        R("Nối next rồi đổi head", 0.50, ["O_C_SLL_INSERT_HEAD"],
          "Thiếu newNode.next = head; head = newNode.",
          "Đổi head trước khi nối next.",
          fn="FUN_INSERT_HEAD"),
        R("RETURN head", 0.25, ["O_C_HEAD_POINTER"],
          "Thiếu RETURN head mới.",
          "Không trả về head sau chèn."),
    ],
    "Q_C3_R_APP_H": [
        R("Rỗng trả null", 0.30, ["O_C_SINGLY_LINKED_LIST"],
          "Thiếu IF head = null RETURN null.",
          "Xóa khi rỗng không an toàn.",
          fn="FUN_IS_EMPTY"),
        R("head = head.next", 0.45, ["O_C_SLL_DELETE_HEAD"],
          "Thiếu head = head.next.",
          "Không cập nhật head.",
          fn="FUN_DELETE_HEAD"),
        R("RETURN head", 0.25, ["O_C_HEAD_POINTER"],
          "Thiếu RETURN head mới.",
          "Không trả về danh sách sau xóa.",
          rule="rul-000031"),
    ],
    "Q_C3_U_DES_E": [
        R("Truy cập thứ i", 0.35, ["O_C_ARRAY", "O_C_LINKED_LIST"],
          "Thiếu mảng O(1) vs DSLK O(n) khi truy cập phần tử thứ i.",
          "So sánh ngược truy cập.",
          fn="FUN_COMPARE_ARRAY_AND_LIST"),
        R("Chèn xóa đầu", 0.35, ["O_C_SLL_INSERT_HEAD", "O_C_ARRAY_INSERT"],
          "Thiếu mảng chèn/xóa đầu O(n) vs DSLK O(1).",
          "Cho mảng chèn đầu O(1).",
          rule="rul-000030"),
        R("Kích thước tĩnh vs động", 0.30, ["O_C_STATIC_STRUCTURE", "O_C_DYNAMIC_STRUCTURE"],
          "Thiếu mảng kích thước cố định, DSLK cấp phát từng node.",
          "Nhầm tĩnh/động.",
          fn="FUN_CAN_GROW"),
    ],
    "Q_C3_U_DES_M": [
        R("Head là cửa vào", 0.40, ["O_C_HEAD_POINTER"],
          "Thiếu head là cửa vào; mất head thì mất danh sách; head=null là rỗng.",
          "Coi head chỉ là node dữ liệu thông thường.",
          rule="rul-000017", fn="FUN_HAS_HEAD"),
        R("Tail chèn cuối O(1)", 0.40, ["O_C_TAIL_POINTER", "O_C_SLL_INSERT_TAIL"],
          "Thiếu tail cho phép gắn node cuối O(1).",
          "Nêu chèn cuối SLL luôn O(1) dù không có tail.",
          rule="rul-000022", fn="FUN_HAS_TAIL"),
        R("Không có tail phải duyệt", 0.20, ["O_C_LINEAR_TIME"],
          "Thiếu ý không có tail thì chèn cuối O(n).",
          "Bỏ qua chi phí tìm node cuối."),
    ],
    "Q_C3_U_DES_H": [
        R("Next cuối trỏ đầu", 0.40, ["O_C_CIRCULAR_LINKED_LIST"],
          "Thiếu next của node cuối trỏ lại node đầu, không null.",
          "Mô tả circular như SLL có next=null.",
          rule="rul-000040", fn="FUN_IS_CIRCULAR"),
        R("Dừng khi quay lại head", 0.35, ["O_C_CIRCULAR_LINKED_LIST"],
          "Thiếu duyệt dừng khi p.next quay lại head, không chờ p=null.",
          "Duyệt đến null trên danh sách vòng (lặp vô hạn).",
          rule="rul-000020"),
        R("Mô hình vòng", 0.25, ["O_C_CIRCULAR_LINKED_LIST"],
          "Thiếu ý phù hợp dữ liệu vòng, từ phần tử hiện tại luôn có next.",
          "Không nêu lợi ích so với SLL thường."),
    ],
    "Q_C3_U_PRO_E": [
        R("Tạo node next null", 0.25, ["O_C_NODE"],
          "Thiếu tạo node mới next = null.",
          "Không khởi tạo next."),
        R("Duyệt tới node cuối", 0.45, ["O_C_SLL_INSERT_TAIL"],
          "Thiếu duyệt đến khi p.next = null rồi gắn node mới.",
          "Chèn ngay tại head.",
          fn="FUN_INSERT_TAIL"),
        R("O(n) vì không tail", 0.30, ["O_C_LINEAR_TIME", "O_C_TAIL_POINTER"],
          "Thiếu giải thích phải duyệt vì không dùng tail, O(n).",
          "Nêu chèn cuối SLL không tail là O(1).",
          rule="rul-000022"),
    ],
    "Q_C3_U_PRO_M": [
        R("Trường hợp rỗng/một node", 0.25, ["O_C_SLL_DELETE_TAIL"],
          "Thiếu xử lý rỗng hoặc chỉ một node (head = null).",
          "Xóa đuôi khi rỗng.",
          fn="FUN_IS_EMPTY"),
        R("Hai con trỏ prev curr", 0.45, ["O_C_SLL_DELETE_TAIL", "O_C_NEXT_POINTER"],
          "Thiếu prev và curr để tìm node cuối trên SLL.",
          "Chỉ dùng một con trỏ nên không cắt được next.",
          fn="FUN_DELETE_TAIL"),
        R("prev.next = null", 0.30, ["O_C_LINEAR_TIME"],
          "Thiếu cắt curr bằng prev.next = null; chi phí O(n).",
          "Không cập nhật liên kết hoặc nêu O(1)."),
    ],
    "Q_C3_U_PRO_H": [
        R("Xuất phát từ tail", 0.35, ["O_C_DLL_TRAVERSE", "O_C_TAIL_POINTER"],
          "Thiếu p = tail rồi đi p = p.prev.",
          "Duyệt từ head về trước trên SLL.",
          fn="FUN_TRAVERSE_LIST"),
        R("Dừng khi p = null", 0.25, ["O_C_POINTER"],
          "Thiếu dừng khi p = null.",
          "Không có điều kiện dừng.",
          fn="FUN_IS_NULL_POINTER"),
        R("SLL không duyệt ngược", 0.40, ["O_C_SINGLY_LINKED_LIST", "O_C_DOUBLY_LINKED_LIST"],
          "Thiếu so sánh SLL không có prev nên không duyệt ngược được.",
          "Cho SLL duyệt ngược như DLL.",
          rule="rul-000033", fn="FUN_SUPPORTS_BACKWARD"),
    ],
    "Q_C3_U_APP_E": [
        R("Duyệt p từ head", 0.30, ["O_C_SLL_SEARCH", "O_C_HEAD_POINTER"],
          "Thiếu p = head và WHILE p <> null.",
          "Truy cập theo chỉ số.",
          fn="FUN_SEARCH_NODE"),
        R("So sánh p.data", 0.45, ["O_C_NODE"],
          "Thiếu IF p.data = x RETURN 1.",
          "Không so data hoặc trả sai.",
          fn="FUN_HAS_DATA"),
        R("Không thấy trả 0", 0.25, ["O_C_OUTPUT"],
          "Thiếu RETURN 0 khi hết danh sách.",
          "Không có nhánh thất bại.",
          rule="rul-000032"),
    ],
    "Q_C3_U_APP_M": [
        R("Danh sách rỗng", 0.30, ["O_C_SLL_INSERT_TAIL"],
          "Thiếu nếu head = null thì head = tail = newNode.",
          "Chèn tail khi rỗng mà không gán head.",
          fn="FUN_IS_EMPTY"),
        R("Gắn vào tail", 0.45, ["O_C_TAIL_POINTER", "O_C_SLL_INSERT_TAIL"],
          "Thiếu tail.next = newNode; tail = newNode.",
          "Duyệt từ head dù đã có tail.",
          rule="rul-000022", fn="FUN_INSERT_TAIL"),
        R("newNode.next = null", 0.25, ["O_C_NODE"],
          "Thiếu newNode.next = null.",
          "Node mới không kết thúc danh sách."),
    ],
    "Q_C3_U_APP_H": [
        R("Ba con trỏ prev curr next", 0.40, ["O_C_SLL_REVERSE"],
          "Thiếu prev, curr, nextTemp khi đảo.",
          "Đảo bằng mảng phụ (không phải ý chính đảo tại chỗ).",
          fn="FUN_REVERSE_LIST"),
        R("Đảo next", 0.40, ["O_C_NEXT_POINTER"],
          "Thiếu curr.next = prev rồi tiến prev/curr.",
          "Gán next sai làm mất node."),
        R("head = prev", 0.20, ["O_C_HEAD_POINTER"],
          "Thiếu RETURN prev như head mới.",
          "Không cập nhật head sau đảo."),
    ],
    "Q_C3_AP_DES_E": [
        R("Chọn DSLK động", 0.40, ["O_C_LINKED_LIST", "O_C_DYNAMIC_STRUCTURE"],
          "Thiếu chọn DSLK vì chèn/xóa đầu thường xuyên.",
          "Chọn mảng cho nhu cầu chèn/xóa đầu.",
          rule="rul-000036", fn="FUN_RECOMMEND_DATA_STRUCTURE"),
        R("Chèn đầu O(1)", 0.35, ["O_C_SLL_INSERT_HEAD", "O_C_CONSTANT_TIME"],
          "Thiếu lập luận chỉ sửa head, O(1).",
          "Nêu chèn đầu DSLK là O(n).",
          rule="rul-000030", fn="FUN_INSERT_HEAD"),
        R("Mảng phải dịch O(n)", 0.25, ["O_C_ARRAY_INSERT", "O_C_STATIC_STRUCTURE"],
          "Thiếu mảng chèn đầu O(n) và kích thước cố định.",
          "Nêu mảng chèn đầu O(1)."),
    ],
    "Q_C3_AP_DES_M": [
        R("Chọn DSLK đôi có tail", 0.40, ["O_C_DOUBLY_LINKED_LIST", "O_C_TAIL_POINTER"],
          "Thiếu chọn DLL có tail cho duyệt ngược và xóa cuối.",
          "Chọn SLL khi cần duyệt ngược.",
          fn="FUN_RECOMMEND_DATA_STRUCTURE"),
        R("Xóa cuối O(1)", 0.35, ["O_C_TAIL_POINTER", "O_C_CONSTANT_TIME"],
          "Thiếu xóa cuối DLL: tail = tail.prev, O(1).",
          "Nêu xóa cuối DLL là O(n)."),
        R("SLL không prev", 0.25, ["O_C_SINGLY_LINKED_LIST"],
          "Thiếu SLL không duyệt ngược; xóa cuối O(n).",
          "Cho SLL xóa cuối O(1) không tail.",
          rule="rul-000033", fn="FUN_SUPPORTS_BACKWARD"),
    ],
    "Q_C3_AP_DES_H": [
        R("Chọn danh sách vòng", 0.45, ["O_C_CIRCULAR_LINKED_LIST"],
          "Thiếu chọn DSLK vòng vì next cuối nối về đầu.",
          "Chọn SLL thường cho mô hình vòng.",
          rule="rul-000040", fn="FUN_IS_CIRCULAR"),
        R("Luôn có next", 0.30, ["O_C_NEXT_POINTER"],
          "Thiếu từ phần tử hiện tại luôn đi tiếp được.",
          "Còn next = null như SLL."),
        R("Tránh lặp vô hạn", 0.25, ["O_C_CIRCULAR_LINKED_LIST"],
          "Thiếu điều kiện dừng khi quay lại điểm xuất phát.",
          "Duyệt đến null trên circular.",
          rule="rul-000020"),
    ],
    "Q_C3_AP_PRO_E": [
        R("Tạo node 9", 0.25, ["O_C_NODE"],
          "Thiếu tạo node 9.",
          "Không tạo node mới."),
        R("Nối vào head", 0.50, ["O_C_SLL_INSERT_HEAD"],
          "Thiếu 9.next = head rồi head = 9.",
          "Cập nhật head trước khi nối.",
          fn="FUN_INSERT_HEAD"),
        R("Kết quả dãy", 0.25, ["O_C_SINGLY_LINKED_LIST"],
          "Thiếu kết quả [9]→[3]→[5]→[8]→null.",
          "Dãy sau chèn sai."),
    ],
    "Q_C3_AP_PRO_M": [
        R("Tìm node trước", 0.40, ["O_C_SLL_DELETE_AT"],
          "Thiếu prev trỏ node trước node cần xóa.",
          "Xóa mà không giữ node trước.",
          fn="FUN_DELETE_NODE"),
        R("Nối prev.next", 0.40, ["O_C_NEXT_POINTER"],
          "Thiếu prev.next = curr.next để cắt node 5.",
          "Mất phần đuôi danh sách."),
        R("Kết quả [3]→[8]", 0.20, ["O_C_SINGLY_LINKED_LIST"],
          "Thiếu kết quả [3]→[8]→null.",
          "Danh sách sau xóa sai."),
    ],
    "Q_C3_AP_PRO_H": [
        R("Khởi tạo prev curr", 0.25, ["O_C_SLL_REVERSE"],
          "Thiếu prev = null, curr = head.",
          "Khởi tạo sai cặp con trỏ.",
          fn="FUN_REVERSE_LIST"),
        R("Đảo từng next", 0.50, ["O_C_NEXT_POINTER", "O_C_SLL_REVERSE"],
          "Thiếu lần lượt đảo next: lưu nextTemp, gán curr.next = prev.",
          "Mất liên kết khi đảo."),
        R("Head mới là prev", 0.25, ["O_C_HEAD_POINTER"],
          "Thiếu head mới là node 3 sau khi curr = null.",
          "Không đổi head sau đảo."),
    ],
    "Q_C3_AP_APP_E": [
        R("WHILE đếm node", 0.45, ["O_C_SLL_COUNT"],
          "Thiếu duyệt p và tăng count đến khi p = null.",
          "Dùng length như mảng.",
          fn="FUN_COUNT_NODES"),
        R("Rỗng cho 0", 0.25, ["O_C_SINGLY_LINKED_LIST"],
          "Thiếu count = 0 khi head = null.",
          "Không xử lý danh sách rỗng.",
          rule="rul-000021", fn="FUN_IS_EMPTY"),
        R("RETURN count", 0.30, ["O_C_OUTPUT"],
          "Thiếu RETURN count.",
          "Không trả về số node."),
    ],
    "Q_C3_AP_APP_M": [
        R("index = 0 như insertHead", 0.35, ["O_C_SLL_INSERT_HEAD", "O_C_SLL_INSERT_AT"],
          "Thiếu nhánh index = 0 chèn đầu.",
          "Không xử lý chèn tại 0.",
          fn="FUN_INSERT_HEAD"),
        R("Tìm node index-1", 0.40, ["O_C_SLL_INSERT_AT"],
          "Thiếu duyệt tới node trước vị trí chèn rồi nối newNode.",
          "Chèn sai chỗ hoặc mất đuôi.",
          fn="FUN_INSERT_AT"),
        R("RETURN head", 0.25, ["O_C_HEAD_POINTER"],
          "Thiếu RETURN head (có thể đã đổi khi chèn đầu).",
          "Không trả về head."),
    ],
    "Q_C3_AP_APP_H": [
        R("Rỗng hoặc index < 0", 0.25, ["O_C_SLL_DELETE_AT"],
          "Thiếu trả về head không đổi khi rỗng hoặc index < 0.",
          "Xóa khi input không hợp lệ.",
          fn="FUN_IS_EMPTY"),
        R("index = 0 xóa đầu", 0.30, ["O_C_SLL_DELETE_HEAD"],
          "Thiếu nhánh index = 0: head = head.next.",
          "Không xử lý xóa đầu.",
          fn="FUN_DELETE_HEAD"),
        R("Cắt p.next", 0.45, ["O_C_SLL_DELETE_AT", "O_C_NEXT_POINTER"],
          "Thiếu tìm node index-1 rồi p.next = p.next.next; kiểm tra next tồn tại.",
          "Cắt sai liên kết hoặc không kiểm tra index vượt độ dài.",
          fn="FUN_DELETE_NODE"),
    ],
}


def load_json(name):
    with open(os.path.join(ONTO, name), "r", encoding="utf-8") as f:
        return json.load(f)


def save_json(name, data):
    path = os.path.join(ONTO, name)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)
        f.write("\n")


def attr_list(pairs):
    return [{"name": k, "value": v} for k, v in pairs if v is not None and v != ""]


def assertion(source, relation, target, attributes=None):
    item = {"source": source, "relation": relation, "target": target}
    if attributes:
        item["attributes"] = attributes
    return item


def build():
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Question"]
    headers = [c.value for c in ws[1]]
    questions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        if d.get("ID"):
            questions.append(d)
    wb.close()

    missing = [q["ID"] for q in questions if q["ID"] not in MODELS]
    extra = [qid for qid in MODELS if qid not in {q["ID"] for q in questions}]
    if missing or extra:
        raise SystemExit(f"Mapping lệch Excel. Thiếu: {missing}. Thừa: {extra}")

    concept_ids = {c["id"] for c in load_json("concepts.json")["concepts"]}
    rule_ids = {r["id"] for r in load_json("rules.json")["rules"]}
    fn_ids = {f["id"] for f in load_json("functions.json")["functions"]}

    instances = []
    assertions = []
    weight_errors = []
    unknown = []

    for q in questions:
        qid = q["ID"]
        rules = MODELS[qid]
        total = round(sum(r["weight"] for r in rules), 10)
        if abs(total - 1.0) > 1e-9:
            weight_errors.append(f"{qid}: {total}")

        q_inst = f"I_{qid}"
        ea_inst = f"I_EA_{qid}"
        instances.append({
            "id": q_inst,
            "instanceOf": "O_C_QUESTION",
            "attributes": attr_list([
                ("name", qid),
                ("content", q["Nội dung câu hỏi"] or ""),
                ("isRubricValid", True),
            ]),
        })
        instances.append({
            "id": ea_inst,
            "instanceOf": "O_C_EXPECTED_ANSWER",
            "attributes": attr_list([
                ("name", f"Expected answer of {qid}"),
                ("description", q["Câu trả lời mong đợi"] or ""),
                ("isExpectedRuleValid", True),
            ]),
        })
        assertions.extend([
            assertion(q_inst, REL_TYPE, q["QuestionTypeID"]),
            assertion(q_inst, REL_DIFF, q["DifficultyID"]),
            assertion(q_inst, REL_BLOOM, q["BloomID"]),
            assertion(q_inst, REL_EA, ea_inst),
        ])
        for rub, w in RUBRICS:
            assertions.append(assertion(q_inst, REL_RUBRIC, rub, [
                {"name": "weight", "value": w}
            ]))

        for i, rule in enumerate(rules, 1):
            er_id = f"I_ER_{qid}_{i}"
            attrs = [
                ("name", rule["name"]),
                ("weight", rule["weight"]),
                ("missExplanation", rule["miss"]),
                ("wrongExplanation", rule["wrong"]),
            ]
            if rule.get("domainRuleId"):
                attrs.append(("domainRuleId", rule["domainRuleId"]))
                if rule["domainRuleId"] not in rule_ids:
                    unknown.append(f"{er_id} rule {rule['domainRuleId']}")
            if rule.get("functionId"):
                attrs.append(("functionId", rule["functionId"]))
                if rule["functionId"] not in fn_ids:
                    unknown.append(f"{er_id} fn {rule['functionId']}")
            instances.append({
                "id": er_id,
                "instanceOf": "O_C_EXPECTED_RULE",
                "attributes": attr_list(attrs),
            })
            assertions.append(assertion(ea_inst, REL_SCORE, er_id))
            for cid in rule["concepts"]:
                if cid not in concept_ids:
                    unknown.append(f"{er_id} concept {cid}")
                assertions.append(assertion(er_id, REL_CONCEPT, cid))

    if weight_errors:
        raise SystemExit("Tổng weight ≠ 1:\n" + "\n".join(weight_errors))
    if unknown:
        raise SystemExit("ID không tồn tại:\n" + "\n".join(unknown))

    save_json("instances.json", {"instances": instances})

    existing = load_json("assertions.json")
    existing["assertions"] = [
        a for a in existing["assertions"]
        if not (
            str(a.get("source", "")).startswith("I_Q_")
            or str(a.get("source", "")).startswith("I_EA_")
            or str(a.get("source", "")).startswith("I_ER_")
        )
    ]
    existing["assertions"].extend(assertions)
    save_json("assertions.json", existing)
    return questions, instances, assertions


def fmt_attrs(attrs):
    if not attrs:
        return None
    lines = []
    for a in attrs:
        lines.append(f"{a['name']}: {a['value']}")
    return "\n".join(lines)


def sync_excel(questions, instances, assertions):
    wb = load_workbook(XLSX)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="334155")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    def style_header(ws):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")

    # Instance sheet
    ws = wb["Instance"]
    ws.delete_rows(2, ws.max_row)
    ws.cell(1, 1, "#")
    ws.cell(1, 2, "ID")
    ws.cell(1, 3, "instanceOf")
    ws.cell(1, 4, "attribute")
    style_header(ws)
    for i, inst in enumerate(instances, 1):
        ws.cell(i + 1, 1, i)
        ws.cell(i + 1, 2, inst["id"])
        ws.cell(i + 1, 3, inst["instanceOf"])
        cell = ws.cell(i + 1, 4, fmt_attrs(inst.get("attributes")))
        cell.alignment = wrap
        for col in range(1, 5):
            ws.cell(i + 1, col).border = thin

    # Assertion: rebuild from JSON to keep DSA + assessment together
    data = load_json("assertions.json")
    ws = wb["Assertion"]
    ws.delete_rows(2, ws.max_row)
    for i, a in enumerate(data["assertions"], 1):
        ws.cell(i + 1, 1, i)
        ws.cell(i + 1, 2, a["source"])
        ws.cell(i + 1, 3, a["relation"])
        ws.cell(i + 1, 4, a["target"])
        attrs = a.get("attributes") or []
        ws.cell(i + 1, 5, fmt_attrs(attrs) if attrs else None)
        for col in range(1, 6):
            ws.cell(i + 1, col).border = thin
            ws.cell(i + 1, col).alignment = wrap

    # Concept row updates
    ws = wb["Concept"]
    headers = [c.value for c in ws[1]]
    id_col = headers.index("ID") + 1
    attr_col = headers.index("attribute") + 1
    name_col = headers.index("name") + 1
    domain_col = headers.index("Domain") + 1
    found_ea = None
    max_num = 0
    for row in range(2, ws.max_row + 1):
        cid = ws.cell(row, id_col).value
        num = ws.cell(row, 1).value
        if isinstance(num, int):
            max_num = max(max_num, num)
        if cid == "O_C_EXPECTED_ANSWER":
            found_ea = row
            ws.cell(row, attr_col).value = (
                'name: str (required, default="Expected Answer")\n'
                "description: str\n"
                "isExpectedRuleValid: bool (required, default=true)"
            )
            ws.cell(row, attr_col).alignment = wrap
    insert_at = (found_ea + 1) if found_ea else ws.max_row + 1
    already = any(ws.cell(r, id_col).value == "O_C_EXPECTED_RULE" for r in range(2, ws.max_row + 1))
    if not already:
        ws.insert_rows(insert_at)
        ws.cell(insert_at, 1, max_num + 1)
        ws.cell(insert_at, id_col, "O_C_EXPECTED_RULE")
        ws.cell(insert_at, name_col, "Expected Rule")
        ws.cell(insert_at, domain_col, "Assessment")
        ws.cell(insert_at, attr_col, (
            'name: str (required, default="Expected Rule")\n'
            "weight: float (required, Domain [0,1])\n"
            "missExplanation: str (required)\n"
            "wrongExplanation: str (required)\n"
            "domainRuleId: str\n"
            "functionId: str"
        ))
        ws.cell(insert_at, attr_col).alignment = wrap

    # Relation rows
    ws = wb["Relation"]
    rel_ids = {ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    new_rels = [
        (None, "REL_EXPECTED_ANSWER_HAS_SCORING_RULE", "hasScoringRule", "O_C_EXPECTED_ANSWER",
         "(1..*)", "O_C_EXPECTED_RULE", None, None,
         "Lời giải mong đợi gồm các luật chấm; tổng weight = 1"),
        (None, "REL_EXPECTED_RULE_REQUIRES_CONCEPT", "requiresConcept", "O_C_EXPECTED_RULE",
         "(0..*)", "Algorithm|Data|... (các lớp DSA)", None, None,
         "Luật chấm yêu cầu các khái niệm DSA tương ứng"),
    ]
    max_num = 0
    for row in range(2, ws.max_row + 1):
        num = ws.cell(row, 1).value
        if isinstance(num, int):
            max_num = max(max_num, num)
    for rel in new_rels:
        if rel[1] in rel_ids:
            continue
        max_num += 1
        r = ws.max_row + 1
        ws.cell(r, 1, max_num)
        for i, val in enumerate(rel[1:], 2):
            ws.cell(r, i, val)

    # Function
    ws = wb["Function"]
    fids = {ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    if "FUN_EXPECTED_RULE_WEIGHT_SUM" not in fids:
        r = ws.max_row + 1
        nums = [ws.cell(i, 1).value for i in range(2, ws.max_row + 1) if isinstance(ws.cell(i, 1).value, int)]
        ws.cell(r, 1, (max(nums) if nums else 0) + 1)
        ws.cell(r, 2, "FUN_EXPECTED_RULE_WEIGHT_SUM")
        ws.cell(r, 3, "expectedRuleWeightSum")
        ws.cell(r, 4, "expectedAnswer: O_C_EXPECTED_ANSWER")
        ws.cell(r, 5, "total: float")
        ws.cell(r, 6, "Tổng trọng số các luật chấm gắn với một lời giải mong đợi")

    # Rule
    ws = wb["Rule"]
    rids = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    if "rul-000053" not in rids:
        r = ws.max_row + 1
        ws.cell(r, 1, "rul-000053")
        ws.cell(r, 2, "expectedRuleWeightValidate")
        ws.cell(r, 3, "expectedRuleWeightSum(ea) = 1.0")
        ws.cell(r, 4, "(attribute) ea.isExpectedRuleValid = True")
        ws.cell(r, 5, "Tổng trọng số các luật chấm của một lời giải mong đợi phải bằng 1.0")

    # Operand
    ws = wb["Operand"]
    oids = {ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    extras = [
        ("OP_EXPECTED_ANSWER", "concept", "ea", "O_C_EXPECTED_ANSWER"),
        ("OP_EXPECTED_RULE_WEIGHT_SUM", "function", "expectedRuleWeightSum()", "FUN_EXPECTED_RULE_WEIGHT_SUM"),
    ]
    nums = [ws.cell(i, 1).value for i in range(2, ws.max_row + 1) if isinstance(ws.cell(i, 1).value, int)]
    n = max(nums) if nums else 0
    for oid, ot, var, val in extras:
        if oid in oids:
            continue
        n += 1
        r = ws.max_row + 1
        ws.cell(r, 1, n)
        ws.cell(r, 2, oid)
        ws.cell(r, 3, ot)
        ws.cell(r, 4, var)
        ws.cell(r, 5, val)

    # ExpectedRule sheet for inspection
    if "ExpectedRule" in wb.sheetnames:
        del wb["ExpectedRule"]
    ws = wb.create_sheet("ExpectedRule")
    cols = ["#", "QuestionID", "ExpectedAnswerID", "ExpectedRuleID", "name", "weight",
            "concepts", "domainRuleId", "functionId", "missExplanation", "wrongExplanation"]
    for i, h in enumerate(cols, 1):
        ws.cell(1, i, h)
    style_header(ws)
    row_i = 2
    n = 1
    er_by_id = {i["id"]: i for i in instances if i["instanceOf"] == "O_C_EXPECTED_RULE"}
    for q in questions:
        qid = q["ID"]
        ea = f"I_EA_{qid}"
        for i, rule in enumerate(MODELS[qid], 1):
            er_id = f"I_ER_{qid}_{i}"
            ws.cell(row_i, 1, n)
            ws.cell(row_i, 2, qid)
            ws.cell(row_i, 3, ea)
            ws.cell(row_i, 4, er_id)
            ws.cell(row_i, 5, rule["name"])
            ws.cell(row_i, 6, rule["weight"])
            ws.cell(row_i, 7, ", ".join(rule["concepts"]))
            ws.cell(row_i, 8, rule.get("domainRuleId"))
            ws.cell(row_i, 9, rule.get("functionId"))
            ws.cell(row_i, 10, rule["miss"])
            ws.cell(row_i, 11, rule["wrong"])
            for col in range(1, 12):
                ws.cell(row_i, col).alignment = wrap
                ws.cell(row_i, col).border = thin
            n += 1
            row_i += 1
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["E"].width = 28
    ws.column_dimensions["G"].width = 40
    ws.column_dimensions["J"].width = 50
    ws.column_dimensions["K"].width = 50

    wb.save(XLSX)


def main():
    questions, instances, assertions = build()
    sync_excel(questions, instances, assertions)
    n_q = len(questions)
    n_er = sum(len(MODELS[q["ID"]]) for q in questions)
    print(f"Questions: {n_q}")
    print(f"Instances: {len(instances)} (Q={n_q}, EA={n_q}, ER={n_er})")
    print(f"New assertions: {len(assertions)}")
    print("Excel synced (Instance, Assertion, ExpectedRule, T-Box rows).")


if __name__ == "__main__":
    main()
